from pathlib import Path
import io,json,os,shutil,subprocess,tempfile,uuid,zipfile,xml.etree.ElementTree as ET
from PySide6.QtCore import QThread,Signal
try: import fitz
except ImportError: fitz=None

class ConversionWorker(QThread):
    progress=Signal(int)
    finished=Signal(bool,str,str)
    def __init__(self,source,target_kind): super().__init__(); self.source=Path(source); self.target_kind=target_kind
    def _progress(self,value): self.progress.emit(max(0,min(100,int(value))))
    def run(self):
        temporary_target = None
        try:
            self._progress(2)
            if not self.source.is_file():
                raise FileNotFoundError(self.source)
            source_suffix = self.source.suffix.lower()
            if self.target_kind not in {"pdf", "docx", "word"}:
                raise ValueError("Unsupported conversion target.")
            target_suffix = ".pdf" if self.target_kind == "pdf" else ".docx"
            if source_suffix == target_suffix:
                raise ValueError("The file is already in the requested format.")
            if (source_suffix, target_suffix) not in {(".docx", ".pdf"), (".pdf", ".docx")}:
                raise ValueError("Only PDF and DOCX files can be converted in place.")
            target = self.source.with_suffix(target_suffix)
            if target.exists():
                raise FileExistsError(
                    "Conversion cancelled because the destination already exists: "
                    + str(target)
                )

            # Build and validate a complete sibling file before changing either final path.
            temporary_target = target.with_name(
                ".{}.digi-tmp-{}{}".format(target.stem, uuid.uuid4().hex, target.suffix)
            )
            self._progress(5)
            if target_suffix == ".docx":
                self.pdf_to_docx(self.source, temporary_target)
            else:
                self.docx_to_pdf(self.source, temporary_target)
            self._validate_output(temporary_target)
            self._progress(95)

            # Final rename sequence:
            # 1) os.link installs the validated file at the new name atomically and
            #    fails rather than overwriting a destination created during conversion.
            # 2) unlink the old name only after the new name is safely installed.
            # 3) remove the temporary name. All paths are in the same directory.
            #
            # This is not a single atomic two-name swap. A crash between steps 1 and 2
            # can leave both names, but the original data remains available.
            try:
                os.link(temporary_target, target)
            except FileExistsError:
                raise FileExistsError(
                    "Conversion cancelled because the destination already exists: "
                    + str(target)
                )
            except OSError as exc:
                raise RuntimeError(
                    "The filesystem could not safely install the converted file "
                    "without overwriting an existing destination. The original was preserved."
                ) from exc

            try:
                self.source.unlink()
            except OSError as exc:
                # Roll back only the output we just installed, never an unrelated file.
                rollback_succeeded = False
                try:
                    target_stat = target.stat()
                    temp_stat = temporary_target.stat()
                    if (target_stat.st_dev, target_stat.st_ino) == (temp_stat.st_dev, temp_stat.st_ino):
                        target.unlink()
                        rollback_succeeded = True
                except OSError:
                    pass
                if rollback_succeeded:
                    message = (
                        "Digi could not remove the original, so conversion was cancelled. "
                        "The original remains unchanged; check file locks and permissions."
                    )
                else:
                    message = (
                        "Digi could not remove the original or safely roll back the new file. "
                        "The original remains, and both filenames may exist. Check permissions "
                        "and inspect the destination before trying again."
                    )
                raise RuntimeError(message) from exc

            try:
                temporary_target.unlink(missing_ok=True)
                temporary_target = None
            except OSError:
                # The final file is installed; the finally block retries temp cleanup.
                pass
            self._progress(100)
            self.finished.emit(True, "Conversion completed successfully.", str(target))
        except Exception as exc:
            self.finished.emit(False, str(exc), "")
        finally:
            if temporary_target is not None:
                try:
                    temporary_target.unlink(missing_ok=True)
                except OSError:
                    pass

    @staticmethod
    def _validate_output(path):
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError("Conversion produced no usable output.")
        if path.suffix.lower() == ".pdf":
            if fitz is None:
                raise RuntimeError("Cannot validate the PDF because PyMuPDF is not installed.")
            with fitz.open(str(path)) as document:
                if len(document) < 1:
                    raise RuntimeError("Conversion produced a PDF with no pages.")
        elif path.suffix.lower() == ".docx":
            with zipfile.ZipFile(path, "r") as document:
                if document.testzip() is not None:
                    raise RuntimeError("Conversion produced a damaged Word document.")
                ET.fromstring(document.read("word/document.xml"))
        else:
            raise RuntimeError("Conversion produced an unsupported output format.")
    def pdf_to_docx(self,source,target):
        if fitz is None: raise RuntimeError("PyMuPDF is not installed.")
        from docx import Document
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_SECTION
        pdf=fitz.open(str(source)); records=[]
        try:
            total=max(1,len(pdf))
            for number,page in enumerate(pdf,1):
                pix=page.get_pixmap(matrix=fitz.Matrix(200/72,200/72),alpha=False)
                records.append({"page":number,"width_pt":float(page.rect.width),"height_pt":float(page.rect.height),"png":pix.tobytes("png")})
                self._progress(5 + (number/total)*55)
        finally: pdf.close()
        if not records: raise RuntimeError("The PDF contains no pages.")
        doc=Document()
        total=len(records)
        for i,r in enumerate(records):
            s=doc.sections[0] if i==0 else doc.add_section(WD_SECTION.NEW_PAGE)
            w,h=r["width_pt"]/72,r["height_pt"]/72
            s.page_width,s.page_height=Inches(w),Inches(h)
            s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(0)
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=p.paragraph_format.space_after=Inches(0)
            p.add_run().add_picture(io.BytesIO(r["png"]),width=Inches(w),height=Inches(h))
            self._progress(60 + ((i+1)/total)*25)
        doc.save(str(target))
        self._progress(88)
        # Store Digi's round-trip metadata as a valid OPC package part.
        # Every ZIP part must have a declared content type; adding a bare JSON
        # part without a [Content_Types].xml declaration makes Word report repair.
        marker={"format":"digi-page-image-docx-v1","dpi":200,"pages":[{"page":r["page"],"width_pt":r["width_pt"],"height_pt":r["height_pt"],"media":f"word/media/image{i+1}.png"} for i,r in enumerate(records)]}
        temp=target.with_name(".{}.package-{}.tmp".format(target.name, uuid.uuid4().hex))
        try:
            with zipfile.ZipFile(target,"r") as zin,zipfile.ZipFile(temp,"w",zipfile.ZIP_DEFLATED) as zout:
                items=zin.infolist()
                total_items=max(1,len(items))
                for i,item in enumerate(items,1):
                    data=zin.read(item.filename)
                    if item.filename=="[Content_Types].xml":
                        # Add a default content type for the JSON metadata part,
                        # preserving the existing XML and namespace declarations.
                        text=data.decode("utf-8")
                        if "Extension=\"json\"" not in text:
                            close=text.rfind("</Types>")
                            if close < 0:
                                raise RuntimeError("DOCX package has an invalid [Content_Types].xml file.")
                            declaration='<Default Extension="json" ContentType="application/json"/>'
                            text=text[:close]+declaration+text[close:]
                            data=text.encode("utf-8")
                    zout.writestr(item,data)
                    self._progress(88 + (i/total_items)*10)
                zout.writestr("word/digi_search_engine_page_images.json",json.dumps(marker).encode("utf-8"))
            os.replace(temp,target)
        finally:
            try: temp.unlink(missing_ok=True)
            except OSError: pass
    def docx_to_pdf(self,source,target):
        if fitz is None: raise RuntimeError("PyMuPDF is not installed.")
        try:
            with zipfile.ZipFile(source,"r") as zin:
                marker="word/digi_search_engine_page_images.json"
                if marker in zin.namelist():
                    info=json.loads(zin.read(marker).decode()); pages=info.get("pages",[]); total=max(1,len(pages)); pdf=fitz.open()
                    self._progress(10)
                    for i,r in enumerate(pages,1):
                        png=zin.read(r["media"]); page=pdf.new_page(width=float(r["width_pt"]),height=float(r["height_pt"]))
                        page.insert_image(fitz.Rect(0,0,page.rect.width,page.rect.height),stream=png)
                        self._progress(10 + (i/total)*80)
                    pdf.save(str(target),deflate=True,clean=True); pdf.close(); self._progress(95); return
        except Exception: pass
        try:
            import pythoncom,win32com.client
            pythoncom.CoInitialize()
            try:
                word=win32com.client.DispatchEx("Word.Application"); word.Visible=False
                document=word.Documents.Open(str(source.resolve()))
                try: document.ExportAsFixedFormat(str(target.resolve()),17,OpenAfterExport=False,OptimizeFor=0,CreateBookmarks=0)
                finally: document.Close(False); word.Quit()
                self._progress(92)
            finally: pythoncom.CoUninitialize()
            return
        except Exception: pass
        libre=shutil.which("soffice") or shutil.which("libreoffice")
        if not libre: raise RuntimeError("Word-to-PDF requires Microsoft Word or LibreOffice.")
        self._progress(20)
        with tempfile.TemporaryDirectory(prefix=".digi-convert-", dir=str(target.parent)) as staging:
            proc=subprocess.run([libre,"--headless","--convert-to","pdf","--outdir",staging,str(source)],capture_output=True,text=True,timeout=120)
            if proc.returncode: raise RuntimeError(proc.stderr or proc.stdout or "LibreOffice conversion failed.")
            generated=Path(staging)/(source.stem+".pdf")
            if not generated.is_file() or generated.stat().st_size == 0:
                raise RuntimeError("LibreOffice did not produce a usable PDF.")
            os.replace(generated,target)
        self._progress(92)
