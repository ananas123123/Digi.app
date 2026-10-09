import json
import time
import urllib.request
from pathlib import Path
from PySide6.QtCore import QObject, Signal, Slot, QThread
from PySide6.QtWidgets import QFileDialog
from .config import APP_VERSION, DEFAULT_INCOMING, LIBRARY_CONFIG, get_library_root, ensure_directories
from .database import Database
from .search import SearchService
from .files import FileService
from .conversion import ConversionWorker
from .incoming import IncomingService
from .notes import NotesService
from .version_manager import initialize_version_file, version_integrity, recalibrate_version


class ReleaseManifestWorker(QThread):
    resultReady = Signal(str)
    URLS = (
        "https://api.github.com/repos/ananas123123/digiwebversionreleases/contents/latest.json?ref=main",
        "https://raw.githubusercontent.com/ananas123123/digiwebversionreleases/main/latest.json",
    )

    def run(self):
        cache_buster = str(int(time.time() * 1000))
        for base_url in self.URLS:
            try:
                separator = "&" if "?" in base_url else "?"
                url = base_url + separator + "_digi_check=" + cache_buster
                request = urllib.request.Request(url, headers={
                    "Accept": "application/vnd.github+json",
                    "User-Agent": "Digi-Update-Checker",
                    "Cache-Control": "no-cache, no-store, max-age=0",
                    "Pragma": "no-cache",
                })
                with urllib.request.urlopen(request, timeout=8) as response:
                    payload = response.read()
                if "api.github.com" in base_url:
                    import base64
                    envelope = json.loads(payload.decode("utf-8"))
                    if envelope.get("encoding") != "base64" or not envelope.get("content"):
                        raise ValueError("GitHub API did not return base64 file content")
                    payload = base64.b64decode(envelope["content"])
                manifest = json.loads(payload.decode("utf-8"))
                self.resultReady.emit(json.dumps({"ok": True, "manifest": manifest}))
                return
            except Exception:
                continue
        self.resultReady.emit(json.dumps({"ok": False, "manifest": None}))


class DigiBridge(QObject):
    @Slot()
    def minimizeWindow(self):
        if self.parent():
            self.parent().showMinimized()

    @Slot()
    def toggleMaximizeWindow(self):
        window = self.parent()
        if window:
            window.showNormal() if window.isMaximized() else window.showMaximized()

    @Slot()
    def closeWindow(self):
        if self.parent():
            self.parent().close()

    indexUpdated = Signal()
    releaseManifestResult = Signal(str)
    conversionProgress = Signal(str, int)
    conversionFinished = Signal(bool, str, str)
    error = Signal(str)

    def __init__(self, parent=None):
        super().__init__(parent)
        initialize_version_file()
        self.version_ok, self.version_problem = version_integrity()
        self.db = None
        self.search_service = None
        self.incoming = None
        self.notes = None
        self.worker = None
        self.release_worker = None
        if self.version_ok:
            try:
                ensure_directories()
                self._initialize_services()
            except Exception as exc:
                self.version_ok = False
                self.version_problem = f"initialization_failed: {exc}"
                self.error.emit(self.version_problem)

    def _initialize_services(self):
        """Create backend services once the runtime layout passes integrity checks."""
        self.db = Database()
        root = get_library_root()
        saved = self.db.setting("incoming_folder")
        incoming = Path(saved) if saved else DEFAULT_INCOMING
        self.search_service = SearchService()
        self.search_service.configure(root, incoming)
        self.incoming = IncomingService(self.db, root, incoming)
        self.notes = NotesService(root)
        self.start_scan()

    @Slot()
    def checkReleaseManifest(self):
        if self.release_worker and self.release_worker.isRunning():
            return
        self.release_worker = ReleaseManifestWorker(self)
        self.release_worker.resultReady.connect(self.releaseManifestResult.emit)
        self.release_worker.start()

    @Slot(result=str)
    def state(self):
        return json.dumps({
            "library": str(self.search_service.root) if self.search_service else "",
            "incoming": str(self.incoming.incoming_folder) if self.incoming else "",
            "version": APP_VERSION,
            "version_ok": self.version_ok,
            "version_problem": self.version_problem,
            "ready": bool(self.db and self.search_service and self.incoming and self.notes),
        })

    @Slot(result=bool)
    def recalibrateVersion(self):
        # Historical method name retained for frontend compatibility. This is
        # now a read-only recheck; it never repairs or rewrites user data.
        self.version_ok = recalibrate_version()
        if self.version_ok:
            self.version_problem = "ok"
            try:
                ensure_directories()
                self._initialize_services()
            except Exception as exc:
                self.version_ok = False
                self.version_problem = f"initialization_failed: {exc}"
                self.error.emit(self.version_problem)
        else:
            self.version_ok, self.version_problem = version_integrity()
        return self.version_ok

    @Slot(str, str, str, str, str, str, result=str)
    def search(self, q, typ, status, source, method, sort):
        if not self.search_service:
            return json.dumps([])
        return json.dumps(self.search_service.query(q, typ, status, source, method, sort))

    @Slot(str, result=str)
    def tags(self, path):
        return json.dumps(self.db.tags(path)) if self.db else "[]"

    @Slot(str, str, result=bool)
    def addTag(self, path, name):
        if not self.db:
            return False
        self.db.add_tag(path, name)
        return True

    @Slot(str, str, result=bool)
    def removeTag(self, path, name):
        if not self.db:
            return False
        self.db.remove_tag(path, name)
        return True

    @Slot(str, str, result=bool)
    def setStatus(self, path, value):
        if not self.db:
            return False
        self.db.update(path, "status", value)
        return True

    @Slot(str, str, result=bool)
    def setDueDate(self, path, value):
        if not self.db:
            return False
        self.db.update(path, "due_date", value)
        return True

    @Slot(str, result=str)
    def previewFile(self, path):
        """Return a safe, read-only preview payload for a selected PDF or Word file."""
        import base64
        from html import escape
        source = Path(path)
        try:
            if not source.is_file():
                return json.dumps({"kind": "error", "message": "This file is no longer available."})
            suffix = source.suffix.lower()
            if suffix == ".pdf":
                try:
                    import fitz
                except ImportError:
                    return json.dumps({"kind": "error", "message": "PDF preview is unavailable because the PDF renderer is not installed."})
                pages = []
                with fitz.open(str(source)) as document:
                    total = len(document)
                    for index, page in enumerate(document):
                        if index >= 20:
                            break
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.15, 1.15), alpha=False)
                        encoded = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
                        pages.append({"number": index + 1, "image": "data:image/png;base64," + encoded})
                return json.dumps({"kind": "pdf", "name": source.name, "pages": pages, "total": total, "truncated": total > len(pages)})
            if suffix in (".docx", ".docm"):
                try:
                    from docx import Document
                    document = Document(str(source))
                    blocks = []
                    for paragraph in document.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            blocks.append("<p>" + escape(text) + "</p>")
                    for table in document.tables:
                        rows = []
                        for row in table.rows:
                            cells = "".join("<td>" + escape(cell.text.strip()) + "</td>" for cell in row.cells)
                            rows.append("<tr>" + cells + "</tr>")
                        blocks.append("<table>" + "".join(rows) + "</table>")
                    body = "".join(blocks) or "<p>This Word document contains no extractable text. Use Open to view it in Word.</p>"
                    return json.dumps({"kind": "word", "name": source.name, "html": body})
                except Exception as exc:
                    return json.dumps({"kind": "error", "message": "Could not read this Word document (" + type(exc).__name__ + "). Use Open to view it in Word."})
            if suffix == ".doc":
                try:
                    import pythoncom
                    import win32com.client
                    pythoncom.CoInitialize()
                    word = None
                    document = None
                    try:
                        word = win32com.client.DispatchEx("Word.Application")
                        word.Visible = False
                        word.DisplayAlerts = 0
                        document = word.Documents.Open(str(source.resolve()), ReadOnly=True, AddToRecentFiles=False)
                        blocks = []
                        for paragraph in document.Paragraphs:
                            text = paragraph.Range.Text.strip("\r\x07\n ")
                            if text:
                                blocks.append("<p>" + escape(text) + "</p>")
                        body = "".join(blocks) or "<p>This Word document contains no extractable text. Use Open to view it in Word.</p>"
                        return json.dumps({"kind": "word", "name": source.name, "html": body})
                    finally:
                        if document is not None:
                            document.Close(False)
                        if word is not None:
                            word.Quit()
                        pythoncom.CoUninitialize()
                except Exception as exc:
                    return json.dumps({"kind": "error", "message": "Previewing legacy .doc files requires Microsoft Word. Details: " + type(exc).__name__ + ". You can still use Open to view the file."})
            return json.dumps({"kind": "error", "message": "Preview is available for PDF and Word documents only."})
        except Exception as exc:
            return json.dumps({"kind": "error", "message": "Preview could not be loaded: " + str(exc)})

    @Slot(str, result=bool)
    def openFile(self, path):
        if not self.incoming:
            return False
        self.incoming.remember_origin(path)
        FileService.open_file(path)
        return True

    @Slot(str, result=bool)
    def openFolder(self, path):
        FileService.open_folder(path)
        return True

    @Slot(result=str)
    def chooseLibraryFolder(self):
        if not self.search_service:
            return ""
        return QFileDialog.getExistingDirectory(
            None, "Choose Digi Search Engine library folder", str(self.search_service.root)
        )

    @Slot(str, result=bool)
    def setLibraryFolder(self, path):
        if not all((self.search_service, self.incoming, self.db)):
            return False
        root = Path(path).resolve()
        self.search_service.configure(root, self.incoming.incoming_folder)
        self.db.set_setting("library_folder", str(root))
        LIBRARY_CONFIG.write_text(str(root), encoding="utf-8")
        self.incoming.set_folders(root, self.incoming.incoming_folder)
        self.notes = NotesService(root)
        self.start_scan()
        return True

    @Slot(result=str)
    def chooseIncomingFolder(self):
        if not all((self.search_service, self.incoming, self.db)):
            return ""
        folder = QFileDialog.getExistingDirectory(
            None, "Choose Incoming folder", str(self.incoming.incoming_folder)
        )
        if folder:
            self.incoming.set_folders(self.search_service.root, folder)
            self.search_service.incoming = Path(folder).resolve()
            self.db.set_setting("incoming_folder", folder)
        return folder or ""

    @Slot(result=bool)
    def processIncoming(self):
        if not self.incoming:
            return False
        self.incoming.process()
        self.start_scan()
        return True

    @Slot(str, str, str, result=str)
    def createDocument(self, parent, name, kind):
        if not self.search_service:
            return ""
        result = FileService.create_document(parent, name, kind)
        self.start_scan()
        return result

    @Slot(str, str, result=str)
    def createFolder(self, parent, name):
        if not self.search_service:
            return ""
        result = FileService.create_folder(parent, name)
        self.start_scan()
        return result

    @Slot(str, result=bool)
    def deleteFolder(self, path):
        if not self.search_service:
            return False
        FileService.delete_folder(path)
        self.start_scan()
        return True

    @Slot(str, str, result=bool)
    def convert(self, path, target):
        if self.worker and self.worker.isRunning():
            return False
        self.worker = ConversionWorker(path, target)
        self.worker.progress.connect(lambda value: self.conversionProgress.emit(str(path), value))
        self.worker.finished.connect(
            lambda ok, msg, out: self.conversionFinished.emit(ok, msg, str(path))
        )
        self.worker.start()
        return True

    @Slot(result=str)
    def notesTree(self):
        return json.dumps(self.notes.tree()) if self.notes else "[]"

    @Slot(str, str, result=str)
    def createNotebook(self, parent, name):
        return self.notes.create_notebook(name) if self.notes else ""

    @Slot(str, str, result=str)
    def createNoteFolder(self, parent, name):
        return self.notes.create_folder(name) if self.notes else ""

    @Slot(str, str, result=str)
    def createNotePage(self, parent, name):
        return self.notes.create_page(name) if self.notes else ""

    @Slot(str, str, result=str)
    def saveNote(self, path, data):
        return self.notes.save(path, data) if self.notes else ""

    @Slot(str, result=str)
    def loadNote(self, path):
        return self.notes.load(path) if self.notes else ""

    @Slot(str, result=bool)
    def deleteNote(self, path):
        if not self.notes:
            return False
        self.notes.delete(path)
        return True

    @Slot(str, str, result=str)
    def renameNote(self, path, name):
        return self.notes.rename(path, name) if self.notes else ""

    @Slot(result=str)
    def recentSearches(self):
        return json.dumps(self.db.recents()) if self.db else "[]"

    @Slot(str, result=bool)
    def rememberSearch(self, q):
        if not self.db:
            return False
        self.db.add_recent(q)
        return True

    @Slot(result=bool)
    def scan(self):
        if not self.search_service:
            return False
        self.start_scan()
        return True

    def start_scan(self):
        if self.search_service:
            self.search_service.scan(
                lambda: self.indexUpdated.emit(),
                lambda message: self.error.emit(message),
            )
